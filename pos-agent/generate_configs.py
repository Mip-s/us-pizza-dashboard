#!/usr/bin/env python3
"""
Generate one Telegraf config per outlet for the POS Uptime Monitor.

Inputs (keep these OUT of git — the tokens are secrets):
  --excel    "Outlet Info - FeedMe.xlsx"   (equipment listing: KDS / kiosk IPs)
  --tokens   heartbeat_tokens.csv          (outlet_code, outlet_name, heartbeat_token)
  --url      https://<your-dashboard-domain>
  --process  POS app process name, e.g. "PosApp.exe" (optional; leave empty to skip the app check)
  --out      output folder (default: ./configs)

Usage:
  pip install openpyxl
  python generate_configs.py --excel "Outlet Info - FeedMe.xlsx" --tokens heartbeat_tokens.csv \
      --url https://us-pizza-dashboard.pages.dev --process PosApp.exe
"""
import argparse
import csv
import os
import re
from collections import defaultdict

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))


def clean(v):
    if v is None:
        return None
    s = re.sub(r"\s+", " ", str(v)).strip()
    return None if s in ("", "-") else s


def load_devices(path):
    """outlet_code -> list of (device_type, ip)"""
    wb = openpyxl.load_workbook(path, data_only=True)
    devices = defaultdict(list)
    for sheet in ("Malaysia Equipment Listing", "Singapore Equipment Listing"):
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        header = [clean(c) or "" for c in next(ws.iter_rows(values_only=True))]
        col = {name.lower(): i for i, name in enumerate(header)}
        i_code, i_type, i_ip = col["outlet code"], col["device type"], col["ip address"]
        for row in list(ws.iter_rows(values_only=True))[1:]:
            code = clean(row[i_code])
            if not code:
                continue
            devices[code].append((clean(row[i_type]), clean(row[i_ip])))
    return devices


def ping_blocks(devs):
    """Numbered KDS-n / SOK-n / ODS-n entries, matching the stations in Supabase."""
    blocks, kds_n, sok_n, ods_n = [], 0, 0, 0
    for dtype, ip in devs:
        if not dtype:
            continue
        if re.fullmatch(r"KDS\d", dtype):
            kds_n += 1
            channel, station = "kds", f"KDS-{kds_n}"
        elif dtype == "Kiosk":
            sok_n += 1
            channel, station = "kiosk", f"SOK-{sok_n}"
        elif re.fullmatch(r"ODS\s*\d*|Order Display.*", dtype, re.I):
            # ODS = order display screen (TV) — must match an "ods" station in Supabase outlet_stations
            ods_n += 1
            channel, station = "ods", f"ODS-{ods_n}"
        else:
            continue
        if ip:
            blocks.append(
                f'[[inputs.ping]]\n'
                f'  urls = ["{ip}"]\n'
                f'  method = "native"\n'
                f'  count = 3\n'
                f'  ping_interval = 1.0\n'
                f'  fieldinclude = ["result_code", "percent_packet_loss"]\n'
                f'  [inputs.ping.tags]\n'
                f'    peer_channel = "{channel}"\n'
                f'    peer_station = "{station}"'
            )
        else:
            blocks.append(
                f'# {station}: IP address not in the equipment listing yet — fill it in and uncomment\n'
                f'# [[inputs.ping]]\n'
                f'#   urls = ["192.168.x.x"]\n'
                f'#   method = "native"\n'
                f'#   count = 3\n'
                f'#   ping_interval = 1.0\n'
                f'#   fieldinclude = ["result_code", "percent_packet_loss"]\n'
                f'#   [inputs.ping.tags]\n'
                f'#     peer_channel = "{channel}"\n'
                f'#     peer_station = "{station}"'
            )
    return "\n\n".join(blocks) if blocks else "# No KDS / kiosk / ODS listed for this outlet"


def procstat_block(process):
    if process:
        return f'[[inputs.procstat]]\n  pid_finder = "native"   # Windows has no pgrep\n  exe = "{process}"\n  fieldinclude = ["running"]'
    return (
        '# POS app check disabled — set the process name and uncomment:\n'
        '# [[inputs.procstat]]\n'
        '#   pid_finder = "native"\n'
        '#   exe = "PosApp.exe"\n'
        '#   fieldinclude = ["running"]'
    )


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--excel", required=True)
    ap.add_argument("--tokens", required=True)
    ap.add_argument("--url", required=True)
    ap.add_argument("--process", default="")
    ap.add_argument("--out", default="configs")
    args = ap.parse_args()

    template = open(os.path.join(HERE, "telegraf.conf.template"), encoding="utf-8").read()
    devices = load_devices(args.excel)
    os.makedirs(args.out, exist_ok=True)

    count = 0
    with open(args.tokens, newline="", encoding="utf-8") as f:
        for row in csv.DictReader(f):
            code = row["outlet_code"]
            conf = (
                template.replace("{{OUTLET_CODE}}", code)
                .replace("{{OUTLET_NAME}}", row["outlet_name"])
                .replace("{{POS_STATION}}", "POS-1")
                .replace("{{DASHBOARD_URL}}", args.url.rstrip("/"))
                .replace("{{TOKEN}}", row["heartbeat_token"])
                .replace("{{PROCSTAT_BLOCK}}", procstat_block(args.process))
                .replace("{{PING_BLOCKS}}", ping_blocks(devices.get(code, [])))
            )
            folder = os.path.join(args.out, code)
            os.makedirs(folder, exist_ok=True)
            with open(os.path.join(folder, "telegraf.conf"), "w", encoding="utf-8", newline="\r\n") as out:
                out.write(conf)
            count += 1
    print(f"Wrote {count} configs to {args.out}/<outlet_code>/telegraf.conf")


if __name__ == "__main__":
    main()
